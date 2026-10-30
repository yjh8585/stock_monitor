#!/usr/bin/env python3
"""MarkLines 중국 엑셀 3종(참고/oem 중국/) → oem_china_* 테이블 적재.

| 파일                          | 테이블                      | 범위                   |
| ----------------------------- | --------------------------- | ---------------------- |
| china_monthly_prod_en.xls     | oem_china_production_month  | 시트 2018~ (MIN_YEAR)  |
| china_exp_en.xlsx             | oem_china_export_month      | 시트 2018~             |
| china_passenger_YYYY.xlsx     | oem_china_segment_month     | 파일 1개 = 1년         |

다운로드는 sync_oem_china_excel.py 가 맡고 이 스크립트는 로컬 파일만 읽는다. 파싱 규칙(승용만·
그룹명 맞추기·n.a.→브랜드·Grand Total 자기검증)은 lib/china_excel.py.

재적재 정책(결정 8): upsert 뒤 **파일이 덮는 기간 안에서 파일에 없는 키를 지운다** — MarkLines 가
모델명을 바꾸면 옛 이름 행이 남아 이중으로 세지기 때문이다. 단 지울 행이 기존의 STALE_ABORT_RATIO 를
넘으면 파싱이 깨진 것으로 보고 **upsert 도 하지 않고** 멈춘다(빈 파일 하나로 전 기간이 지워지거나,
새 키만 들어가 옛 키와 이중으로 세지는 사고 방지). 판정은 lib/china_excel.stale_keys.
upsert_rows 가 캐시 무효화를 자동 수행하고, 삭제는 WriteSession 이 무효화한다.

사용법
-----
  python scripts/import_oem_china.py              # 로컬 파일 전부 적재
  python scripts/import_oem_china.py --dry-run    # DB 없이 행 수·교차 검증만
  python scripts/import_oem_china.py --only nev   # nev | export | passenger
경위·결정 = docs/plans/2026-10-29-oem-china-tab.md
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl
import xlrd
from loguru import logger

sys.path.insert(0, str(Path(__file__).parent))
from lib.bootstrap import init_script  # noqa: E402

init_script(__file__)

from lib.china_excel import (  # noqa: E402
  ParseResult,
  export_sheet_name,
  parse_export_sheet,
  parse_nev_sheet,
  parse_passenger_sheet,
  stale_keys,
)
from lib.db import WriteSession, get_client, upsert_rows  # noqa: E402

CHINA_DIR = Path(__file__).resolve().parents[1] / '참고' / 'oem 중국'
NEV_FILE = CHINA_DIR / 'china_monthly_prod_en.xls'
EXPORT_FILE = CHINA_DIR / 'china_exp_en.xlsx'
PASSENGER_RE = re.compile(r'china_passenger_(\d{4})\.xlsx$')

#: 적재 시작 연도 — NEV 파워트레인 열이 2018 시트부터 있다(결정 5).
MIN_YEAR = 2018
#: 지울 행이 기존 행의 이 비율을 넘으면 파싱 고장으로 보고 멈춘다.
STALE_ABORT_RATIO = 0.05

TABLES = {
  'nev': ('oem_china_production_month', ('cn_group', 'maker_brand', 'vehicle_type', 'model', 'powertrain', 'year_month')),
  'passenger': ('oem_china_segment_month', ('vehicle_type', 'segment', 'year_month')),
  'export': ('oem_china_export_month', ('cn_group', 'maker', 'brand', 'model', 'year_month')),
}


def read_xls_sheets(path: Path) -> dict[str, list[tuple]]:
  """구형 .xls → {시트명: 행 튜플 리스트}. 날짜 셀은 datetime, 빈 셀은 None."""
  book = xlrd.open_workbook(str(path))
  out = {}
  for sh in book.sheets():
    rows = []
    for i in range(sh.nrows):
      row = []
      for j in range(sh.ncols):
        c = sh.cell(i, j)
        if c.ctype == xlrd.XL_CELL_DATE:
          row.append(xlrd.xldate_as_datetime(c.value, book.datemode))
        elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or c.value == '':
          row.append(None)
        else:
          row.append(c.value)
      rows.append(tuple(row))
    out[sh.name] = rows
  return out


def read_xlsx_sheets(path: Path) -> dict[str, list[tuple]]:
  """.xlsx → {시트명: 행 튜플 리스트}.

  ⚠️ read_only=True 로 열면 MarkLines 엑셀이 조용히 0행이 나온 적이 있다(gotchas-data-collection).
  """
  wb = openpyxl.load_workbook(path, data_only=True)
  try:
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
  finally:
    wb.close()


def _merge(results: list[ParseResult]) -> ParseResult:
  merged = ParseResult([], {})
  for r in results:
    merged.rows.extend(r.rows)
    merged.all_types_total.update(r.all_types_total)
  return merged


def _require_years(path: Path, years: list[int], sheetnames) -> None:
  # 시트 이름 규칙이 바뀌면 0행이 「성공」으로 끝나고 파일까지 교체돼 다음 회차가 건너뛴다
  if not years:
    raise ValueError(f'{path.name}: {MIN_YEAR}년 이후 연도 시트를 하나도 못 찾았다 — 시트 이름 {list(sheetnames)}')


def parse_nev_file(path: Path) -> tuple[ParseResult, list[int]]:
  sheets = read_xls_sheets(path)
  years = sorted(int(n) for n in sheets if n.isdigit() and int(n) >= MIN_YEAR)
  _require_years(path, years, sheets)
  return _merge([parse_nev_sheet(y, sheets[str(y)]) for y in years]), years


def parse_export_file(path: Path) -> tuple[ParseResult, list[int]]:
  sheets = read_xlsx_sheets(path)
  years = []
  results = []
  for y in range(MIN_YEAR, 2100):
    name = export_sheet_name(list(sheets), y)
    if name is None:
      continue
    years.append(y)
    results.append(parse_export_sheet(y, sheets[name]))
  _require_years(path, years, sheets)
  return _merge(results), years


def parse_passenger_file(path: Path) -> tuple[ParseResult, list[int]]:
  m = PASSENGER_RE.search(path.name)
  if not m:
    raise ValueError(f'승용차 파일명이 china_passenger_YYYY.xlsx 가 아니다: {path.name}')
  sheets = read_xlsx_sheets(path)
  return parse_passenger_sheet(sheets['Sheet1']), [int(m.group(1))]


def passenger_files() -> list[Path]:
  return sorted(p for p in CHINA_DIR.glob('china_passenger_*.xlsx') if PASSENGER_RE.search(p.name))


def cross_check(nev: ParseResult, passenger: ParseResult) -> list[int]:
  """NEV 와 승용차 파일의 전 차종 월합이 다른 월. 두 파일은 같은 생산이라 원래 같다.

  다운로드 시점이 엇갈리면 최신 월만 다를 수 있어 실패가 아니라 경고로 둔다.
  """
  bad = sorted(
    ym for ym, v in passenger.all_types_total.items()
    if ym in nev.all_types_total and nev.all_types_total[ym] != v
  )
  for ym in bad:
    logger.warning(f'교차 검증 불일치 {ym}: NEV {nev.all_types_total[ym]:,} ≠ 승용차 {passenger.all_types_total[ym]:,}')
  return bad


def _fetch_keys(table: str, pk: tuple[str, ...], ym_from: int, ym_to: int) -> set[tuple]:
  client = get_client()
  keys: set[tuple] = set()
  start = 0
  while True:
    q = client.table(table).select(','.join(pk)).gte('year_month', ym_from).lte('year_month', ym_to)
    for c in pk:
      q = q.order(c)  # 🔴 .range() 다중 페이지는 PK 전체로 정렬해야 행이 안 빠진다
    page = q.range(start, start + 999).execute().data or []
    keys.update(tuple(r[c] for c in pk) for r in page)
    if len(page) < 1000:
      return keys
    start += 1000


def load(
  kind: str, result: ParseResult, years: list[int], dry_run: bool, max_stale_ratio: float = STALE_ABORT_RATIO
) -> int:
  """한 종류 적재 + 같은 기간의 stale 키 삭제. 반환 = upsert 행 수.

  max_stale_ratio 는 표기 규칙을 일부러 바꿔 다시 적재할 때만 올린다(평소엔 기본값).
  """
  table, pk = TABLES[kind]
  by_year: dict[int, int] = {}
  for r in result.rows:
    by_year[r['year_month'] // 100] = by_year.get(r['year_month'] // 100, 0) + 1
  logger.info(f'{table}: {len(result.rows):,}행 · 연도별 {dict(sorted(by_year.items()))}')
  if dry_run:
    return len(result.rows)

  ym_from, ym_to = min(years) * 100 + 1, max(years) * 100 + 12
  # 지울 키 판정을 upsert «전»에 — 멈출 때 DB 에 아무것도 안 들어가 있어야 이중 집계가 안 남는다
  existing = _fetch_keys(table, pk, ym_from, ym_to)
  fresh = {tuple(r[c] for c in pk) for r in result.rows}
  try:
    stale = stale_keys(existing, fresh, max_stale_ratio)
  except RuntimeError as e:
    raise RuntimeError(f'{table}: {e}') from e
  upserted = upsert_rows(table, result.rows, ','.join(pk))
  if stale:
    with WriteSession() as w:
      for key in stale:
        q = w.table(table).delete()
        for c, v in zip(pk, key):
          q = q.eq(c, v)
        q.execute()
    logger.info(f'{table}: 파일에 없는 옛 행 {len(stale):,}개 삭제 ({ym_from}~{ym_to})')
  return upserted


def import_nev(path: Path = NEV_FILE, dry_run: bool = False) -> ParseResult:
  result, years = parse_nev_file(path)
  load('nev', result, years, dry_run)
  return result


def import_export(path: Path = EXPORT_FILE, dry_run: bool = False) -> ParseResult:
  result, years = parse_export_file(path)
  load('export', result, years, dry_run)
  return result


def import_passenger(path: Path, dry_run: bool = False) -> ParseResult:
  result, years = parse_passenger_file(path)
  load('passenger', result, years, dry_run)
  return result


def main(argv: list[str] | None = None) -> int:
  ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
  ap.add_argument('--dry-run', action='store_true', help='DB 없이 파싱·행 수·교차 검증만')
  ap.add_argument('--only', choices=('nev', 'export', 'passenger'))
  args = ap.parse_args(argv)

  nev = None
  if args.only in (None, 'nev'):
    nev = import_nev(dry_run=args.dry_run)
  if args.only in (None, 'export'):
    import_export(dry_run=args.dry_run)
  if args.only in (None, 'passenger'):
    files = passenger_files()
    if not files:
      logger.warning(f'승용차 파일 없음: {CHINA_DIR}/china_passenger_YYYY.xlsx')
    for p in files:
      res = import_passenger(p, dry_run=args.dry_run)
      if nev is not None:
        cross_check(nev, res)
  logger.success('중국 엑셀 적재 완료' + (' (dry-run)' if args.dry_run else ''))
  return 0


if __name__ == '__main__':
  sys.exit(main())
